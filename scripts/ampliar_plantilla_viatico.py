# -*- coding: utf-8 -*-
"""Agrega 3 filas de «Profesional» (2, 3 y 4) al bloque Viático de assets/plantilla_ficha_terreno.xlsx.
Se ejecuta UNA vez; si ya están, no hace nada. (Reutiliza insertar() de ampliar_plantilla_ficha.py.)
Uso: python ampliar_plantilla_viatico.py [plantilla.xlsx]
"""
import sys, copy
from pathlib import Path
import openpyxl
from openpyxl.worksheet.cell_range import CellRange

RAIZ = Path(__file__).resolve().parents[1]
RUTA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "assets" / "plantilla_ficha_terreno.xlsx"


def insertar(ws, idx, n, fila_modelo, extender=()):
    alturas = {r: ws.row_dimensions[r].height for r in range(idx, ws.max_row + 1)}
    merges = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.merged_cells.remove(m)
    ws.insert_rows(idx, n)
    for m in merges:
        if m.min_row >= idx:
            m.shift(row_shift=n)
        elif m.max_row >= idx:
            m.expand(down=n)
        ws.merged_cells.add(m)
    for r in range(ws.max_row, idx - 1, -1):
        ws.row_dimensions[r + n].height = alturas.get(r)
    modelo = fila_modelo if fila_modelo < idx else fila_modelo + n
    for k in range(n):
        r = idx + k
        ws.row_dimensions[r].height = ws.row_dimensions[modelo].height or 22.5
        for c in range(1, 21):
            src = ws.cell(modelo, c)
            if src.has_style:
                ws.cell(r, c)._style = copy.copy(src._style)
    for rango in extender:
        for k in range(n):
            ws.merge_cells(start_row=idx + k, start_column=rango[0], end_row=idx + k, end_column=rango[1])


wb = openpyxl.load_workbook(RUTA)
ws = wb.worksheets[0]
if str(ws["B34"].value or "").startswith("@@VIA_P2"):
    print("Ya está ampliada."); sys.exit(0)
assert ws["B33"].value == "Profesional", ws["B33"].value
insertar(ws, 34, 3, 33)
ws["B33"] = "@@VIA_P1_LBL@@"
for n, fila in ((2, 34), (3, 35), (4, 36)):
    ws.cell(fila, 2).value = f"@@VIA_P{n}_LBL@@"
    for col, k in ((3, "C"), (4, "D"), (5, "E"), (6, "F"), (7, "G")):
        ws.cell(fila, col).value = f"@@VIA_P{n}_{k}@@"
    for col in (15, 16):          # O/P: tarifas auxiliares, no se copian
        ws.cell(fila, col).value = None
if ws.print_area:
    ws.print_area = "B1:P74"
wb.save(RUTA)
print("Plantilla ampliada:", RUTA)
