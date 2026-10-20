# -*- coding: utf-8 -*-
"""Construye assets/plantilla_rendicion_cuentas.xlsx: la planilla "RENDICIÓN DE CUENTAS"
en blanco, con TOKENS @@...@@ en las celdas rellenables (misma estrategia que la ficha de terreno).
Las celdas que en el original eran fórmulas (J17, J20, J30) también quedan como token: la app
calcula los montos y los escribe ya resueltos.

Uso:  python construir_plantilla_rendicion.py <origen.xlsx> [salida.xlsx]
(el origen es cualquier "Ren Ctas ....xlsx" real; la plantilla resultante NO conserva sus datos)
"""
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment

RAIZ = Path(__file__).resolve().parents[1]
SALIDA_DEFAULT = RAIZ / "assets" / "plantilla_rendicion_cuentas.xlsx"
MAX_DOCS = 5  # filas 25..29


def sembrar(ws):
    t = {
        "E5": "@@FUNCIONARIO@@", "H6": "@@MES@@", "B7": "@@ANIO@@",
        "F7": "@@RESOLUCION@@", "H7": "@@FECHA_RESOLUCION@@",
        "F8": "@@IMPUTACION@@", "C9": "@@CHEQUE_NUM@@", "H9": "@@CHEQUE_FECHA@@",
        "E10": "@@TRANSFERENCIA@@",
        "J13": "@@VALOR_RINDE@@", "J15": "@@FONDOS@@", "J17": "@@GASTOS@@",
        "E19": "@@REINTEGRO_DETALLE@@", "J20": "@@REINTEGRO@@",
        "C30": "@@OBSERVACIONES@@", "J30": "@@TOTAL@@", "B36": "@@FECHA_RENDICION@@",
    }
    for i in range(1, MAX_DOCS + 1):
        f = 24 + i
        t[f"B{f}"] = f"@@D{i}_NUM@@"; t[f"C{f}"] = f"@@D{i}_PROV@@"
        t[f"E{f}"] = f"@@D{i}_MAT@@"; t[f"J{f}"] = f"@@D{i}_VAL@@"
    for ref, v in t.items():
        ws[ref] = v
        if ref[0] == "J":
            ws[ref].alignment = Alignment(horizontal="right", vertical="center")


if __name__ == "__main__":
    origen = sys.argv[1]
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else SALIDA_DEFAULT
    wb = openpyxl.load_workbook(origen)
    ws = wb.worksheets[0]
    sembrar(ws)
    # El original trae la FIRMA del funcionario como imagen: se descartan todas menos el logo
    # (el logo es la imagen cuadrada de ~149x135 px; la firma es apaisada).
    ws._images = [im for im in ws._images if abs(im.width / im.height - 1.1) < 0.2]
    wb.properties.creator = "SERNAGEOMIN"
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    print("Plantilla guardada:", salida)
