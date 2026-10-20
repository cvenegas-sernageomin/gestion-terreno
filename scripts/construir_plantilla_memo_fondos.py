# -*- coding: utf-8 -*-
"""Construye assets/plantilla_memo_fondos.docx: el MEMORÁNDUM "solicita autorización de fondos a
rendir sobre 15 UTM" en blanco, con TOKENS @@...@@ (el JS de la app los rellena con JSZip).
Marcadores que la app CLONA: @@JUST@@ (un párrafo con viñeta por punto) y la fila de la tabla
con @@T_MONTO@@ (una por línea de monto).

Uso:  python construir_plantilla_memo_fondos.py <memo_origen.docx> [salida.docx]
(el origen es un memo real, p. ej. "Memo Nº xxx.docx" de abril 2026; la plantilla NO conserva sus datos)
"""
import sys, copy
from pathlib import Path
import docx
from docx.text.paragraph import Paragraph
from docx.table import Table

RAIZ = Path(__file__).resolve().parents[1]
SALIDA_DEFAULT = RAIZ / "assets" / "plantilla_memo_fondos.docx"


def quitar(run):
    run._r.getparent().remove(run._r)


def valor_tras_tab(p, token, desde_tabs=2):
    """Párrafo 'ETIQUETA \t : \t valor': conserva hasta el 2º tab y pone el token en la 1ª run de valor."""
    runs = p.runs
    tabs = 0
    for i, r in enumerate(runs):
        if r.text == "\t":
            tabs += 1
            if tabs == desde_tabs:
                if i + 1 < len(runs):
                    runs[i + 1].text = token
                    for x in runs[i + 2:]:
                        quitar(x)
                return
    raise RuntimeError("no encontré el 2º tab en: " + p.text)


def todo_un_token(p, token):
    runs = p.runs
    runs[0].text = token
    for r in runs[1:]:
        quitar(r)


def construir(origen, salida):
    d = docx.Document(origen)
    body = list(d.element.body)
    P = lambda i: Paragraph(body[i], d)

    # 0 MEMORANDUM N° xx
    r = P(0).runs
    r[2].text = "@@NUM@@"
    for x in r[3:]:
        quitar(x)
    valor_tras_tab(P(2), "@@A_NOMBRE@@")
    todo = P(3).runs; todo[2].text = "@@A_CARGO@@"
    for x in todo[3:]: quitar(x)
    valor_tras_tab(P(5), "@@DE_NOMBRE@@")
    todo = P(6).runs; todo[2].text = "@@DE_CARGO@@"
    for x in todo[3:]: quitar(x)
    valor_tras_tab(P(8), "@@MAT@@")
    valor_tras_tab(P(10), "@@FECHA@@")

    todo_un_token(P(13), "@@P1@@")
    todo_un_token(P(15), "@@P2@@")
    todo_un_token(P(17), "@@P3@@")
    todo_un_token(P(19), "@@JUST@@")
    for i in (20, 21):
        body[i].getparent().remove(body[i])
    todo_un_token(P(23), "@@P4@@")

    # tabla: fila 1 = modelo de línea; se eliminan las filas 2 y 3; total en la fila final
    tbl = Table(body[25], d)
    filas = tbl.rows
    c = filas[1].cells
    for cel, tok in zip(c, ("@@T_FUNC@@", "@@T_DET@@", "@@T_MONTO@@")):
        todo_un_token(cel.paragraphs[0], tok)
        for extra in cel.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
    for fila in list(filas)[2:-1]:
        tbl._tbl.remove(fila._tr)
    todo_un_token(tbl.rows[-1].cells[2].paragraphs[0], "@@T_TOTAL@@")

    # 27: texto antes / monto en negrita / texto después
    p = P(27); rs = p.runs
    rs[0].text = "@@P5A@@"
    rs[2].text = "@@MONTO_TOTAL@@"   # negrita
    rs[4].text = "@@P5B@@"
    for i, x in enumerate(rs):
        if i not in (0, 2, 4):
            quitar(x)

    todo = P(31).runs; todo[0].text = "@@FIRMA_NOMBRE@@"
    for x in todo[1:]: quitar(x)
    todo = P(32).runs; todo[0].text = "@@FIRMA_CARGO@@"
    for x in todo[1:]: quitar(x)
    todo = P(33).runs; todo[0].text = "@@INICIALES@@"
    for x in todo[1:]: quitar(x)

    d.core_properties.author = "SERNAGEOMIN"
    d.core_properties.last_modified_by = "SERNAGEOMIN"
    d.core_properties.title = ""
    salida = Path(salida); salida.parent.mkdir(parents=True, exist_ok=True)
    d.save(salida)
    print("Plantilla guardada:", salida)


if __name__ == "__main__":
    construir(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else SALIDA_DEFAULT)
